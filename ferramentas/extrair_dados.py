"""Extrai os dados dos arquivos originais para o seeder do banco.

- boletos/ (qualquer subpasta) -> server/seed/extraido-multas.json (multas do Detran e da PRF)
  e server/seed/extraido-guias-carro.json (guia de remoção e depósito do Detran).
  O tipo de cada PDF sai do conteúdo, e o arquivo é renomeado e movido para a pasta do tipo:
    boletos/multas/detran/<série>.pdf   boletos/multas/prf/<série>.pdf   boletos/carro/patio-<guia>.pdf
  Então basta soltar o PDF novo em boletos/ e rodar este script.
- despesas/*.xlsx (planilha de despesas; usa a mais recente) -> server/seed/extraido-planilha.json

Os ajustes combinados (o que remover, o que é devolução, custos extras do carro...)
ficam em server/seed/base.js. Depois de extrair, rode `npm run importar` (ou reinicie o
servidor): os boletos novos entram no banco sem apagar nada do que foi lançado no app.

Uso (a partir da pasta do projeto):
    python3 ferramentas/extrair_dados.py

Requer pypdf e openpyxl (pip install pypdf openpyxl).
"""
import glob
import json
import os
import re
import sys
import unicodedata
from datetime import date, datetime, timedelta

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PASTA_BOLETOS = os.path.join(ROOT, "boletos")


def centavos(valor):
    if valor is None or valor == "":
        return None
    if isinstance(valor, str):
        return int(valor.replace(".", "").replace(",", ""))
    return int(round(valor * 100))


def buscar(padrao, texto, grupo=1):
    m = re.search(padrao, texto, re.S)
    return m.group(grupo).strip() if m else None


def slug(texto):
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", texto.lower()).strip("-")


def gravar_json(nome, dados):
    destino = os.path.join(ROOT, "server", "seed", nome)
    with open(destino, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return os.path.relpath(destino, ROOT)


# ---------------------------------------------------------------- Multas


def texto_pdf(caminho):
    from pypdf import PdfReader

    return "\n".join(p.extract_text() for p in PdfReader(caminho).pages)


def mod10(numero):
    total, peso = 0, 2
    for digito in reversed(numero):
        produto = int(digito) * peso
        total += produto // 10 + produto % 10
        peso = 1 if peso == 2 else 2
    return (10 - total % 10) % 10


def mod11_arrecadacao(numero):
    total, peso = 0, 2
    for digito in reversed(numero):
        total += int(digito) * peso
        peso = 2 if peso == 9 else peso + 1
    resto = total % 11
    return 0 if resto in (0, 1) else 11 - resto


def mod11_boleto(numero):
    total, peso = 0, 2
    for digito in reversed(numero):
        total += int(digito) * peso
        peso = 2 if peso == 9 else peso + 1
    dv = 11 - total % 11
    return 1 if dv in (0, 10, 11) else dv


def conferir_arrecadacao(caminho, linha, valor):
    """Guias do Detran (código começa com 8): DVs e valor embutido no código."""
    blocos = linha.split()
    codigo = "".join(b[:11] for b in blocos)
    dv = mod10 if codigo[2] in "67" else mod11_arrecadacao
    if not all(dv(b[:11]) == int(b[11]) for b in blocos) or dv(codigo[:3] + codigo[4:]) != int(codigo[3]):
        sys.exit(f"{caminho}: dígitos verificadores da linha digitável não conferem")
    if int(codigo[4:15]) != valor:
        sys.exit(f"{caminho}: o código de barras traz {int(codigo[4:15])} e o boleto diz {valor}")


def conferir_boleto_bancario(caminho, linha):
    """Boleto bancário (47 dígitos): DVs. Devolve (vencimento, valor) embutidos no código."""
    d = re.sub(r"\D", "", linha)
    campos = [(d[0:9], d[9]), (d[10:20], d[20]), (d[21:31], d[31])]
    if len(d) != 47 or any(mod10(c) != int(v) for c, v in campos):
        sys.exit(f"{caminho}: dígitos verificadores da linha digitável não conferem")
    codigo = d[0:4] + d[32] + d[33:47] + d[4:9] + d[10:20] + d[21:31]
    if mod11_boleto(codigo[:4] + codigo[5:]) != int(codigo[4]):
        sys.exit(f"{caminho}: dígito verificador geral do boleto não confere")
    fator = int(d[33:37])
    # Fator de vencimento: a contagem recomeçou em 1000 no dia 22/02/2025.
    vencimento = date(2025, 2, 22) + timedelta(days=fator - 1000)
    return vencimento.strftime("%d/%m/%Y"), int(d[37:47])


def tipo_documento(t):
    if "GAD/M" in t:
        return "multa-detran"
    if "POLÍCIA RODOVIÁRIA FEDERAL" in t and "CPAG" in t:
        return "multa-prf"
    if "REMOÇÃO E DEPÓSITO" in t:
        return "guia-deposito"
    return None


def ler_multa_detran(t, caminho):
    infracao = re.search(r"Valor\(R\$\)\s*\n(\d{5}) (.+?)(?:\s+CTB|\n)", t)
    amparo = re.search(r"CTB\s+(\d+, ?[IVX]+)\s+(\d+)\s+(\w+)\s+([\d.,]+)", t)
    local = re.search(r"Local\s*\nMunicípio\s*\n(.+?)\n(.+?)\n", t)
    data_hora = re.search(r"Data Hora\s*\n(\d{2}/\d{2}/\d{4}) (\d{2}:\d{2})", t)
    blocos = re.search(r"(\d{12})\s+(\d{12})\s+(\d{12})\s+(\d{12})", t)

    # Guia emitida antes do vencimento da multa: com 20% de desconto. Depois dele: valor cheio com juros.
    cobranca = "com-desconto" if "Valor com 20% já descontado" in t else "com-juros"
    if cobranca == "com-desconto":
        valor = buscar(r"Valor com 20% já descontado\s*\n?R\$\s*([\d.,]+)", t)
        juros = None
    else:
        valor = buscar(r"Valor desta guia\s*\nR\$\s*([\d.,]+)", t)
        juros = buscar(r"Valor a pagar\s*\nR\$\s*[\d.,]+\s*\nR\$\s*([\d.,]+)", t)

    velocidade = None
    if "Velocidade Considerada" in t:
        velocidade = {
            "medida": int(buscar(r"Velocidade Medida\.*:\s*(\d+)", t)),
            "considerada": int(buscar(r"Velocidade Considerada\.*:\s*(\d+)", t)),
            "limite": int(buscar(r"Velocidade Limite\.*:\s*(\d+)", t)),
        }

    boleto = {
        "serie": buscar(r"SérieÓrgão Fiscalizador\s*\n\s*([A-Z]\d{9})", t),
        "placa": buscar(r"Placa UF\s*\n(\w+) \w{2}\n", t),
        "cobranca": cobranca,
        "arquivo": os.path.relpath(caminho, ROOT),
        "codigoInfracao": infracao.group(1),
        "descricao": infracao.group(2).strip(),
        "amparo": "CTB art. " + amparo.group(1),
        "pontos": int(amparo.group(2)),
        "infrator": amparo.group(3).capitalize(),
        "dataInfracao": data_hora.group(1),
        "horaInfracao": data_hora.group(2),
        "local": local.group(1).strip(),
        "municipio": local.group(2).strip(),
        "autuador": buscar(r"\(AUTUADOR\)\s*\n\d+ - (.+?)\n", t),
        "origem": buscar(r"ORIGEM AIT \d+/(\w+)", t),
        "velocidade": velocidade,
        "guia": buscar(r"Nº Guia\s*\n(\d+)", t),
        "expedida": buscar(r"Data Expedição\s*\n(\d{2}/\d{2}/\d{4})", t),
        "validade": buscar(r"Guia Válida Até\s*\n(\d{2}/\d{2}/\d{4})", t),
        "valorOriginal": centavos(amparo.group(4)),
        "juros": centavos(juros) if juros else 0,
        "valor": centavos(valor),
        "linhaDigitavel": " ".join(blocos.groups()),
    }
    faltando = [k for k, v in boleto.items() if v is None and k not in ("origem", "velocidade")]
    if faltando:
        sys.exit(f"{caminho}: não consegui ler {faltando}")
    conferir_arrecadacao(caminho, boleto["linhaDigitavel"], boleto["valor"])
    return boleto


def ler_multa_prf(t, caminho):
    """Boleto GRU da PRF (Consulta e Pagamento de Multas - CPAG)."""
    linhas = [l.strip() for l in t.splitlines()]
    placa = re.search(r"R\$ ([\d.,]+)([A-Z]{3}\d[A-Z0-9]\d{2})\n", t)
    i = linhas.index(f"R$ {placa.group(1)}{placa.group(2)}")
    codigo, descricao = linhas[i + 1], re.sub(r"\s+", " ", linhas[i + 2])
    # O local pode quebrar em várias linhas até a linha com data e hora.
    j = i + 3
    while not re.fullmatch(r"\d{2}/\d{2}/\d{4} \d{2}:\d{2}", linhas[j]):
        j += 1
    local_completo = " ".join(linhas[i + 3 : j])
    local, municipio = [x.strip() for x in local_completo.rsplit(" - ", 1)]
    data_infracao, hora_infracao = linhas[j].split()
    amparo = linhas[j + 1]
    desconto = re.search(r"\n\d{7}\nR\$ ([\d.,]+)\nR\$ ([\d.,]+)\n", t)
    linha_digitavel = re.sub(
        r"\s+", " ", buscar(r"(\d{5}\.\d{5}\s+\d{5}\.\d{6}\s+\d{5}\.\d{6}\s+\d\s+\d{14})", t)
    )

    boleto = {
        "serie": buscar(r"\n([A-Z]\d{9})\n", t),
        "placa": placa.group(2),
        # O boleto da PRF sempre traz o desconto de 20% até o vencimento.
        "cobranca": "com-desconto",
        "arquivo": os.path.relpath(caminho, ROOT),
        "codigoInfracao": codigo,
        "descricao": descricao,
        "amparo": f"CTB {amparo}",
        "pontos": None,
        "infrator": None,
        "dataInfracao": data_infracao,
        "horaInfracao": hora_infracao,
        "local": local,
        "municipio": re.sub(r"/([A-Z]{2})$", r"-\1", municipio),
        "autuador": "PRF - POLÍCIA RODOVIÁRIA FEDERAL",
        "origem": None,
        "velocidade": None,
        "guia": buscar(r"\n(\d{20})\n", t),
        "expedida": buscar(r"Gerado em: (\d{2}/\d{2}/\d{4})", t),
        "validade": buscar(r"Cedente\nPRF[^\n]*\n[^\n]*\n(\d{2}/\d{2}/\d{4})", t),
        "valorOriginal": centavos(placa.group(1)),
        "juros": 0,
        "valor": centavos(desconto.group(2)),
        "linhaDigitavel": linha_digitavel,
    }
    faltando = [k for k, v in boleto.items() if v is None and k not in ("origem", "velocidade", "pontos", "infrator")]
    if faltando:
        sys.exit(f"{caminho}: não consegui ler {faltando}")
    if boleto["valorOriginal"] - centavos(desconto.group(1)) != boleto["valor"]:
        sys.exit(f"{caminho}: valor do documento menos desconto não bate com o valor cobrado")
    vencimento, valor_codigo = conferir_boleto_bancario(caminho, linha_digitavel)
    if vencimento != boleto["validade"] or valor_codigo != boleto["valorOriginal"]:
        sys.exit(f"{caminho}: vencimento ou valor do código de barras não batem com o boleto")
    return boleto


def ler_guia_deposito(t, caminho):
    """Guia de despesas de remoção e depósito de veículos (GAD-E) do Detran."""
    emissao = re.search(r"(\d{2}/\d{2}/\d{4}) (\d{2}:\d{2}) (\d{2}/\d{2}/\d{4})", t)
    itens = [
        {"codigo": m.group(1), "descricao": m.group(2).strip(), "quantidade": int(m.group(4)), "valor": centavos(m.group(5))}
        for m in re.finditer(r"\n(\d{4}) ([A-ZÇÃÁÉÍÓÚÂÊÔ ]+?) (CRD\d+) (\d+) ([\d.]+,\d{2})", t)
    ]
    blocos = re.search(r"(\d{12}) (\d{12}) (\d{12}) (\d{12})", t)
    guia = {
        "tipo": "deposito",
        "guia": buscar(r"\n(\d{10})\n", t),
        "placa": buscar(r"([A-Z]{3}\d[A-Z0-9]\d{2})Placa:", t),
        "arquivo": os.path.relpath(caminho, ROOT),
        "emissao": emissao.group(1) if emissao else None,
        "validade": emissao.group(3) if emissao else None,
        "entrada": buscar(r"Fim\n(\d{2}/\d{2}/\d{4})", t),
        "itens": itens,
        "valor": centavos(buscar(r"R\$ ([\d.]+,\d{2})", t)),
        "linhaDigitavel": " ".join(blocos.groups()) if blocos else None,
    }
    faltando = [k for k, v in guia.items() if v in (None, [])]
    if faltando:
        sys.exit(f"{caminho}: não consegui ler {faltando}")
    if sum(i["valor"] for i in itens) != guia["valor"]:
        sys.exit(f"{caminho}: a soma dos serviços não bate com o total da guia")
    conferir_arrecadacao(caminho, guia["linhaDigitavel"], guia["valor"])
    return guia


def organizar(caminho, destino_relativo):
    """Move o PDF para a pasta e o nome do seu tipo (boletos/multas/detran/N006153327.pdf...)."""
    destino = os.path.join(PASTA_BOLETOS, destino_relativo)
    if os.path.abspath(caminho) == os.path.abspath(destino):
        return caminho
    if os.path.exists(destino):
        sys.exit(f"{os.path.relpath(caminho, ROOT)}: já existe {os.path.relpath(destino, ROOT)} (PDF repetido?)")
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    os.rename(caminho, destino)
    print(f"Organizado: {os.path.relpath(caminho, ROOT)} -> {os.path.relpath(destino, ROOT)}")
    return destino


def apagar_pastas_vazias():
    for pasta, _, _ in sorted(os.walk(PASTA_BOLETOS), key=lambda x: -len(x[0])):
        if pasta != PASTA_BOLETOS and not [n for n in os.listdir(pasta) if n != ".DS_Store"]:
            for n in os.listdir(pasta):
                os.remove(os.path.join(pasta, n))
            os.rmdir(pasta)


def gerar_multas():
    boletos, guias, ignorados = [], [], []
    pdfs = sorted(
        os.path.join(pasta, nome)
        for pasta, _, nomes in os.walk(PASTA_BOLETOS)
        for nome in nomes
        if nome.lower().endswith(".pdf")
    )
    for caminho in pdfs:
        t = texto_pdf(caminho)
        tipo = tipo_documento(t)
        if tipo == "multa-detran":
            serie = buscar(r"SérieÓrgão Fiscalizador\s*\n\s*([A-Z]\d{9})", t)
            boletos.append(ler_multa_detran(t, organizar(caminho, f"multas/detran/{serie}.pdf")))
        elif tipo == "multa-prf":
            serie = buscar(r"\n([A-Z]\d{9})\n", t)
            boletos.append(ler_multa_prf(t, organizar(caminho, f"multas/prf/{serie}.pdf")))
        elif tipo == "guia-deposito":
            guia = buscar(r"\n(\d{10})\n", t)
            guias.append(ler_guia_deposito(t, organizar(caminho, f"carro/patio-{guia}.pdf")))
        else:
            ignorados.append(os.path.relpath(caminho, ROOT))
    apagar_pastas_vazias()
    boletos.sort(key=lambda b: b["serie"])
    guias.sort(key=lambda g: (g["emissao"][6:] + g["emissao"][3:5] + g["emissao"][:2], g["guia"]))
    series = [b["serie"] for b in boletos]
    repetidas = {s for s in series if series.count(s) > 1}
    if repetidas:
        sys.exit(f"Séries repetidas entre os PDFs: {sorted(repetidas)}")
    destino = gravar_json("extraido-multas.json", boletos)
    print(f"{len(boletos)} multas gravadas em {destino}")
    destino = gravar_json("extraido-guias-carro.json", guias)
    print(f"{len(guias)} guia(s) do carro gravada(s) em {destino}")
    for caminho in ignorados:
        print(f"AVISO: formato não reconhecido, ignorado: {caminho}")


# ---------------------------------------------------------------- Planilha


def planilha_mais_recente():
    arquivos = [
        a
        for a in glob.glob(os.path.join(ROOT, "despesas", "*.xlsx"))
        if not os.path.basename(a).startswith("~$")
    ]
    return max(arquivos, key=os.path.getmtime) if arquivos else None


def ler_planilha(caminho):
    import openpyxl

    ws_f = openpyxl.load_workbook(caminho).worksheets[0]
    ws_v = openpyxl.load_workbook(caminho, data_only=True).worksheets[0]

    # Linha de cabeçalho: a primeira com várias datas (os meses).
    for linha_meses in range(1, 11):
        colunas = [
            c
            for c in range(2, ws_v.max_column + 1)
            if isinstance(ws_v.cell(linha_meses, c).value, datetime)
        ]
        if len(colunas) >= 3:
            break
    else:
        sys.exit("Não encontrei a linha com os meses na planilha.")
    meses = [ws_v.cell(linha_meses, c).value.strftime("%Y-%m") for c in colunas]

    def valores(r):
        return [centavos(ws_v.cell(r, c).value) for c in colunas]

    def formulas(r):
        return [
            v
            for c in colunas
            if isinstance(v := ws_f.cell(r, c).value, str) and v.startswith("=")
        ]

    rotulos = {
        r: str(ws_v.cell(r, 1).value).strip()
        for r in range(linha_meses + 1, ws_v.max_row + 1)
        if ws_v.cell(r, 1).value
    }

    def achar(padrao):
        linha = next((r for r, t in rotulos.items() if re.search(padrao, t, re.I)), None)
        if linha is None:
            sys.exit(f"Não encontrei a linha '{padrao}' na planilha.")
        return linha

    # Grupos: linhas cujo valor é =SUM(Xa:Xb) das linhas logo abaixo.
    grupos = {}
    for r in rotulos:
        for f in formulas(r):
            m = re.fullmatch(r"=SUM\([A-Z]+(\d+):[A-Z]+(\d+)\)", f.replace(" ", ""), re.I)
            if m:
                grupos[r] = (int(m.group(1)), int(m.group(2)))
                break

    linha_total = achar(r"^total despesas")
    linha_entradas = achar(r"^entradas")
    linha_saldo = achar(r"saldo.*m[eê]s")
    linha_acumulado = achar(r"saldo.*acumulado")
    linha_aplicacao = achar(r"aplica")
    linha_total_ga = achar(r"^total pgt")

    refs_total = {int(x) for f in formulas(linha_total) for x in re.findall(r"[A-Z]+(\d+)", f)}

    def itens_do_grupo(r_grupo):
        nome_grupo = rotulos[r_grupo]
        inicio, fim = grupos[r_grupo]
        itens = []
        for r in range(inicio, fim + 1):
            if r not in rotulos:
                continue
            nome = rotulos[r]
            itens.append(
                {
                    "id": f"{slug(nome_grupo)}/{slug(nome)}",
                    "nome": nome,
                    "valores": valores(r),
                    "devolucao": bool(re.search(r"devolu", nome, re.I)),
                    "carro": bool(re.search(r"carro", nome, re.I)),
                }
            )
        return itens

    grupos_despesa = [
        {"id": slug(rotulos[r]), "nome": rotulos[r], "itens": itens_do_grupo(r), "totais": valores(r)}
        for r in sorted(grupos)
        if r in refs_total
    ]
    outros = [r for r in sorted(grupos) if r not in refs_total]

    # Meses em que a Greici e o André pagaram as contas (coluna do "Total PGT Greici/André").
    total_ga = valores(linha_total_ga)
    meses_resgate = [meses[i] for i, v in enumerate(total_ga) if v]

    pagos_greici = []
    for r in outros:
        for item in itens_do_grupo(r):
            for i, v in enumerate(item["valores"]):
                if v:
                    pagos_greici.append({"id": item["id"], "nome": item["nome"], "mes": meses[i], "valor": v})
    nome_pagos_greici = rotulos[outros[0]] if outros else ""

    formula_aplicacao = next(iter(formulas(linha_aplicacao)), "")
    taxa = re.search(r"\*\s*([\d.,]+)\s*%", formula_aplicacao)
    aplicacao = valores(linha_aplicacao)
    inicio_aplicacao = next((meses[i] for i, v in enumerate(aplicacao) if v), None)

    dados = {
        "arquivo": os.path.relpath(caminho, ROOT),
        "geradoEm": date.today().strftime("%d/%m/%Y"),
        "meses": meses,
        "mesesResgate": meses_resgate,
        "grupos": grupos_despesa,
        "pagosPelaGreici": {"nome": nome_pagos_greici, "itens": pagos_greici},
        "totalResgate": sum(v for v in total_ga if v),
        "entradas": valores(linha_entradas),
        "totalDespesas": valores(linha_total),
        "saldoMes": valores(linha_saldo),
        "saldoAcumulado": valores(linha_acumulado),
        "aplicacao": {
            "taxaMensal": round(float(taxa.group(1).replace(",", ".")) / 100, 6) if taxa else None,
            "inicio": inicio_aplicacao,
            "valores": aplicacao,
        },
    }
    conferir_planilha(dados)
    return dados


def conferir_planilha(d):
    """Confere se a soma dos itens bate com os totais calculados pela planilha."""
    avisos = []
    for i, mes in enumerate(d["meses"]):
        soma_total = 0
        for g in d["grupos"]:
            soma = sum(item["valores"][i] or 0 for item in g["itens"])
            if soma != (g["totais"][i] or 0):
                avisos.append(f"{mes} {g['nome']}: itens somam {soma}, planilha diz {g['totais'][i]}")
            soma_total += soma
        if soma_total != (d["totalDespesas"][i] or 0):
            avisos.append(f"{mes} total despesas: {soma_total} x {d['totalDespesas'][i]}")
        if d["saldoMes"][i] is not None:
            esperado = (d["entradas"][i] or 0) - soma_total
            if esperado != d["saldoMes"][i]:
                avisos.append(f"{mes} saldo do mês: {esperado} x {d['saldoMes'][i]}")
    resgate = sum(d["totalDespesas"][d["meses"].index(m)] or 0 for m in d["mesesResgate"])
    resgate += sum(item["valor"] for item in d["pagosPelaGreici"]["itens"])
    if resgate != d["totalResgate"]:
        avisos.append(f"total pago pela Greici/André: {resgate} x {d['totalResgate']}")
    for aviso in avisos:
        print("AVISO:", aviso)


def gerar_planilha():
    caminho = planilha_mais_recente()
    if not caminho:
        print("Sem planilha em despesas/: mantendo server/seed/extraido-planilha.json como está.")
        return
    dados = ler_planilha(caminho)
    destino = gravar_json("extraido-planilha.json", dados)
    print(f"Planilha '{dados['arquivo']}' ({len(dados['meses'])} meses) gravada em {destino}")


def main():
    gerar_multas()
    gerar_planilha()


if __name__ == "__main__":
    main()
